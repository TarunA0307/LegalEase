"""
generator.py - Formatting utilities for LegalEase.

Converts the raw AI-generated text into:
  * clean text            -> sanitize_text(text)
  * a Word document       -> format_docx(text, doc_type, ...)   (python-docx)
  * a branded PDF         -> format_pdf(text, doc_type, ...)    (fpdf2)
  * a styled HTML preview -> format_html_preview(text)
  * a plain .txt download -> format_txt(text)

All formatters share one parser (parse_blocks) so the preview, DOCX and PDF
always show the same structure: title, section headings, paragraphs, bullets.
"""

from __future__ import annotations

import html
import io
import os
import re
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config  # noqa: E402

# ---------------------------------------------------------------------------
# 1. Text sanitising
# ---------------------------------------------------------------------------
_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"', "\u201f": '"',
    "\u2032": "'", "\u2033": '"',
    "\u2013": "-", "\u2014": "-", "\u2015": "-", "\u2212": "-", "\u2010": "-", "\u2011": "-",
    "\u2026": "...",
    "\u00a0": " ", "\u2002": " ", "\u2003": " ", "\u2009": " ", "\u202f": " ",
    "\u2022": "-", "\u25cf": "-", "\u25aa": "-", "\u2023": "-", "\u2043": "-",
    "\u20b9": "Rs. ",  # Indian Rupee sign
    "\u2122": "(TM)", "\u00ae": "(R)", "\u00a9": "(C)",
    "\u200b": "", "\u200c": "", "\u200d": "", "\ufeff": "",
}


def sanitize_text(text: str) -> str:
    """Remove typographic quotes/special characters so every format renders cleanly."""
    if not text:
        return ""
    for bad, good in _REPLACEMENTS.items():
        text = text.replace(bad, good)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ")
    # Drop remaining control characters (keep newlines)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Strip code fences if the model wrapped the whole answer in them
    text = re.sub(r"^\s*```[a-zA-Z]*\s*\n", "", text)
    text = re.sub(r"\n\s*```\s*$", "", text)
    # Collapse 3+ blank lines into one blank line
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def format_txt(text: str) -> str:
    """Plain-text export: Markdown symbols removed, structure kept."""
    lines = []
    for block in parse_blocks(text):
        plain = strip_inline(block.text)
        if block.kind == "title":
            lines += [plain.upper(), "=" * min(len(plain), 80), ""]
        elif block.kind == "heading":
            lines += ["", plain, "-" * min(len(plain), 80)]
        elif block.kind == "bullet":
            lines.append(f"  - {plain}")
        elif block.kind == "rule":
            lines.append("-" * 40)
        elif block.kind == "blank":
            if lines and lines[-1] != "":
                lines.append("")
        else:
            lines.append(plain)
    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------
# 2. Parsing the AI output into blocks
# ---------------------------------------------------------------------------
@dataclass
class Block:
    kind: str    # title | heading | subheading | bullet | paragraph | rule | blank
    text: str = ""
    level: int = 0


_BULLET_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_NUMBERED_HEADING_RE = re.compile(r"^(\d+(\.\d+)*\.?|[IVX]+\.|Article\s+\d+[.:]?|Section\s+\d+[.:]?)\s+[^.]{2,70}:?$", re.I)


def parse_blocks(text: str) -> list[Block]:
    """Split document text into typed blocks."""
    blocks: list[Block] = []
    for raw in sanitize_text(text).split("\n"):
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            if blocks and blocks[-1].kind != "blank":
                blocks.append(Block("blank"))
            continue
        if re.fullmatch(r"[-*_=]{3,}", stripped) and not re.fullmatch(r"_{3,}", stripped):
            blocks.append(Block("rule"))
            continue

        m = _HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            content = m.group(2).strip().strip("*").strip()
            has_title = any(b.kind == "title" for b in blocks)
            kind = "title" if level <= 2 and not has_title else "heading"
            blocks.append(Block(kind, content, level))
            continue

        m = _BULLET_RE.match(line)
        if m and not re.fullmatch(r"[-*_ ]+", stripped):
            indent = len(m.group(1).replace("\t", "    ")) // 2
            blocks.append(Block("bullet", m.group(3).strip(), min(indent, 3)))
            continue

        # A whole line in bold, e.g. "**1. Services:**" -> sub-heading
        if re.fullmatch(r"\*\*[^*]+\*\*:?", stripped) and len(stripped) < 100:
            blocks.append(Block("subheading", stripped.replace("**", "").strip()))
            continue

        # Short numbered line such as "1. Services:" -> heading
        if _NUMBERED_HEADING_RE.match(stripped) and stripped.endswith(":") and len(stripped) < 80:
            blocks.append(Block("heading", stripped.rstrip(":").strip(), 3))
            continue

        blocks.append(Block("paragraph", stripped))

    while blocks and blocks[0].kind == "blank":
        blocks.pop(0)
    while blocks and blocks[-1].kind == "blank":
        blocks.pop()
    return blocks


# Inline markdown -> runs of (text, bold, italic)
_INLINE_RE = re.compile(r"(\*\*[^*]+?\*\*|\*[^*\s][^*]*?\*|(?<![A-Za-z0-9_])_[^_\s][^_]*?_(?![A-Za-z0-9_]))")


def inline_runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split a line into (text, bold, italic) runs. Signature lines '____' stay intact."""
    runs = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        token = m.group(0)
        if token.startswith("**"):
            runs.append((token[2:-2], True, False))
        else:
            runs.append((token[1:-1], False, True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    return [r for r in runs if r[0]]


def strip_inline(text: str) -> str:
    return "".join(r[0] for r in inline_runs(text))


def split_terms(terms: str | None) -> list[str]:
    """'a; b; c' -> ['a', 'b', 'c']"""
    if not terms:
        return []
    return [sanitize_text(t).strip() for t in terms.split(";") if t.strip()]


def _document_title(blocks: list[Block], doc_type: str) -> str:
    for b in blocks:
        if b.kind == "title":
            return strip_inline(b.text)
    return doc_type or "Legal Document"


def _logo_stream(logo) -> io.BytesIO | None:
    """Accept a file path, bytes or file-like object and return a PNG stream (or None)."""
    if logo is None:
        logo = config.LOGO_PATH if os.path.exists(config.LOGO_PATH) else None
    if logo is None:
        return None
    try:
        from PIL import Image

        if isinstance(logo, (bytes, bytearray)):
            img = Image.open(io.BytesIO(logo))
        elif hasattr(logo, "read"):
            data = logo.read()
            if hasattr(logo, "seek"):
                logo.seek(0)
            img = Image.open(io.BytesIO(data))
        else:
            img = Image.open(str(logo))
        img.load()
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGBA")
        out = io.BytesIO()
        img.save(out, format="PNG")
        out.seek(0)
        return out
    except Exception:
        return None


def safe_filename(doc_type: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", doc_type or "legal_document").strip("_").lower()
    return f"{base or 'legal_document'}.{ext}"


# ---------------------------------------------------------------------------
# 3. DOCX (python-docx)
# ---------------------------------------------------------------------------
def format_docx(text: str, doc_type: str, terms: str | None = None, logo=None,
                footer_text: str | None = None, include_disclaimer: bool = True) -> bytes:
    """Build a Word document: logo, Times New Roman, headings, bullets, terms table, footer."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    NAVY = RGBColor(0x1F, 0x2A, 0x44)
    blocks = parse_blocks(text)
    title = _document_title(blocks, doc_type)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)

    # Base font: Times New Roman 12pt
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for sname in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        st = doc.styles[sname]
        st.font.name = "Times New Roman"
        st.font.color.rgb = NAVY
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def add_runs(paragraph, content: str, bold_all: bool = False):
        for chunk, bold, italic in inline_runs(content):
            run = paragraph.add_run(chunk)
            run.bold = bold or bold_all
            run.italic = italic

    # Logo
    logo_stream = _logo_stream(logo)
    if logo_stream:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_stream, width=Cm(5))

    # Title
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_after = Pt(14)
    tr = tp.add_run(title.upper())
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.color.rgb = NAVY

    # Body
    for b in blocks:
        if b.kind == "title":
            continue  # already printed
        if b.kind == "blank":
            continue
        if b.kind == "rule":
            doc.add_paragraph("_" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b.kind == "heading":
            h = doc.add_heading(level=2 if b.level <= 3 else 3)
            add_runs(h, b.text)
            for r in h.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13 if b.level <= 3 else 12)
                r.font.color.rgb = NAVY
        elif b.kind == "subheading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            add_runs(p, b.text, bold_all=True)
        elif b.kind == "bullet":
            p = doc.add_paragraph(style="List Bullet" if b.level == 0 else "List Bullet 2")
            add_runs(p, b.text)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, b.text)

    # Terms table (from the semicolon-separated input)
    term_list = split_terms(terms)
    if term_list:
        doc.add_paragraph()
        h = doc.add_heading("Schedule A - Key Terms Summary", level=2)
        for r in h.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.color.rgb = NAVY
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        hdr = table.rows[0].cells
        for cell, label in zip(hdr, ("No.", "Term / Condition")):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:color"), "auto")
            shade.set(qn("w:fill"), "1F2A44")
            cell._tc.get_or_add_tcPr().append(shade)
        for i, term in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = term
        widths = (Cm(1.5), Cm(14.5))
        for col, w in zip(table.columns, widths):
            col.width = w
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = w

    # Disclaimer
    if include_disclaimer:
        doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run(config.DISCLAIMER)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Footer: company text + "Page X"
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run(f"{footer_text or config.FOOTER_TEXT}   |   Page ")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    num_run = OxmlElement("w:r")
    num_props = OxmlElement("w:rPr")
    num_size = OxmlElement("w:sz")
    num_size.set(qn("w:val"), "18")  # 9pt
    num_props.append(num_size)
    num_run.append(num_props)
    num_text = OxmlElement("w:t")
    num_text.text = "1"
    num_run.append(num_text)
    fld.append(num_run)
    footer_p._p.append(fld)

    doc.core_properties.title = title
    doc.core_properties.author = config.COMPANY_NAME
    doc.core_properties.subject = doc_type

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# 4. PDF (fpdf2)
# ---------------------------------------------------------------------------
_UNICODE_FONT_SETS = [
    # Windows
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf",
     "C:/Windows/Fonts/ariali.ttf", "C:/Windows/Fonts/arialbi.ttf"),
    # Linux
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    # macOS
    ("/System/Library/Fonts/Supplemental/Arial.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"),
]


def format_pdf(text: str, doc_type: str, terms: str | None = None, logo=None,
               footer_text: str | None = None, include_disclaimer: bool = True) -> bytes:
    """Build a branded PDF: logo + title header on every page, bold headings,
    bullet-style terms, footer with company text and page numbers."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    blocks = parse_blocks(text)
    title = _document_title(blocks, doc_type)
    logo_stream = _logo_stream(logo)
    logo_bytes = logo_stream.getvalue() if logo_stream else None
    logo_ratio = 0.3  # height / width
    if logo_bytes:
        from PIL import Image

        with Image.open(io.BytesIO(logo_bytes)) as im:
            logo_ratio = im.height / max(im.width, 1)
    footer_line = footer_text or config.FOOTER_TEXT

    class LegalPDF(FPDF):
        font_family_name = "Helvetica"
        unicode_ok = False

        def clean(self, s: str) -> str:
            if self.unicode_ok:
                return s
            return s.encode("latin-1", "replace").decode("latin-1")

        def header(self):
            y = 10
            if logo_bytes:
                # Fit the logo in a 45 x 18 mm box, centred
                w_mm, h_mm = 45.0, 45.0 * logo_ratio
                if h_mm > 18:
                    w_mm, h_mm = 18 / logo_ratio, 18.0
                self.image(io.BytesIO(logo_bytes), x=(self.w - w_mm) / 2, y=y, w=w_mm, h=h_mm)
                self.set_y(y + h_mm + 2)
            else:
                self.set_y(y)
            self.set_font(self.font_family_name, "B", 13)
            self.set_text_color(31, 42, 68)
            self.cell(0, 8, self.clean(title), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_draw_color(180, 180, 180)
            self.line(self.l_margin, self.get_y() + 1, self.w - self.r_margin, self.get_y() + 1)
            self.ln(6)
            self.set_text_color(0, 0, 0)

        def footer(self):
            self.set_y(-15)
            self.set_font(self.font_family_name, "I", 8)
            self.set_text_color(110, 110, 110)
            self.cell(0, 6, self.clean(f"{footer_line}   |   Page {self.page_no()}/{{nb}}"),
                      align="C")

    pdf = LegalPDF(format="A4")
    pdf.set_margins(20, 15, 20)
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.set_title(title)
    pdf.set_author(config.COMPANY_NAME)

    # Use a Unicode TTF font when available, otherwise the built-in Helvetica
    for regular, bold, italic, bold_italic in _UNICODE_FONT_SETS:
        if all(os.path.exists(f) for f in (regular, bold, italic, bold_italic)):
            try:
                pdf.add_font("Body", "", regular)
                pdf.add_font("Body", "B", bold)
                pdf.add_font("Body", "I", italic)
                pdf.add_font("Body", "BI", bold_italic)
                LegalPDF.font_family_name = "Body"
                LegalPDF.unicode_ok = True
                break
            except Exception:
                continue

    fam = LegalPDF.font_family_name
    pdf.add_page()
    body_size, line_h = 10.5, 5.6

    def write_runs(content: str, size: float = body_size, bold_all: bool = False,
                   indent: float = 0.0):
        """Write a line with inline bold/italic, wrapping at the (indented) margin."""
        old_margin = pdf.l_margin
        if indent:
            pdf.set_left_margin(old_margin + indent)
            pdf.set_x(old_margin + indent)
        for chunk, bold, italic in inline_runs(content):
            style = ("B" if (bold or bold_all) else "") + ("I" if italic else "")
            pdf.set_font(fam, style, size)
            pdf.write(line_h, pdf.clean(chunk))
        pdf.ln(line_h)
        if indent:
            pdf.set_left_margin(old_margin)
            pdf.set_x(old_margin)

    for b in blocks:
        if b.kind == "title":
            continue
        if b.kind == "blank":
            pdf.ln(1.5)
        elif b.kind == "rule":
            y = pdf.get_y() + 2
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(5)
        elif b.kind == "heading":
            pdf.ln(2)
            if pdf.get_y() > pdf.h - 40:
                pdf.add_page()
            pdf.set_text_color(31, 42, 68)
            write_runs(b.text, size=12 if b.level <= 3 else 11, bold_all=True)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        elif b.kind == "subheading":
            pdf.ln(1)
            write_runs(b.text, bold_all=True)
        elif b.kind == "bullet":
            indent = 5 + 5 * b.level
            pdf.set_x(pdf.l_margin + indent - 4)
            pdf.set_font(fam, "B", body_size)
            pdf.write(line_h, "-" if not LegalPDF.unicode_ok else "\u2022")
            write_runs(" " + b.text, indent=indent)
        else:
            write_runs(b.text)
            pdf.ln(1)

    # Bullet-style terms summary
    term_list = split_terms(terms)
    if term_list:
        pdf.ln(4)
        pdf.set_text_color(31, 42, 68)
        write_runs("Schedule A - Key Terms Summary", size=12, bold_all=True)
        pdf.set_text_color(0, 0, 0)
        pdf.ln(1)
        for i, term in enumerate(term_list, 1):
            write_runs(f"**{i}.** {term}", indent=5)

    if include_disclaimer:
        pdf.ln(6)
        pdf.set_text_color(100, 100, 100)
        pdf.set_font(fam, "I", 8)
        pdf.multi_cell(0, 4, pdf.clean(config.DISCLAIMER))
        pdf.set_text_color(0, 0, 0)

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# 5. HTML preview (Streamlit)
# ---------------------------------------------------------------------------
def _inline_html(content: str) -> str:
    out = []
    for chunk, bold, italic in inline_runs(content):
        s = html.escape(chunk)
        if bold:
            s = f"<strong>{s}</strong>"
        if italic:
            s = f"<em>{s}</em>"
        out.append(s)
    return "".join(out)


def format_html_preview(text: str) -> str:
    """Convert the document into semantic HTML blocks for the dark preview card."""
    parts: list[str] = []
    in_list = False
    for b in parse_blocks(text):
        if b.kind != "bullet" and in_list:
            parts.append("</ul>")
            in_list = False
        if b.kind == "title":
            parts.append(f"<h2 class='le-title'>{_inline_html(b.text)}</h2>")
        elif b.kind == "heading":
            parts.append(f"<h3 class='le-heading'>{_inline_html(b.text)}</h3>")
        elif b.kind == "subheading":
            parts.append(f"<p class='le-sub'><strong>{_inline_html(b.text)}</strong></p>")
        elif b.kind == "bullet":
            if not in_list:
                parts.append("<ul class='le-list'>")
                in_list = True
            pad = f" style='margin-left:{b.level * 18}px'" if b.level else ""
            parts.append(f"<li{pad}>{_inline_html(b.text)}</li>")
        elif b.kind == "rule":
            parts.append("<hr/>")
        elif b.kind == "paragraph":
            parts.append(f"<p>{_inline_html(b.text)}</p>")
    if in_list:
        parts.append("</ul>")
    return "\n".join(parts)


PREVIEW_CSS = """
<style>
.le-card {background:#111827;border:1px solid #1f2937;border-radius:12px;
  padding:28px 32px;max-height:560px;overflow-y:auto;color:#e5e7eb;
  font-family:Georgia,'Times New Roman',serif;line-height:1.65;font-size:15px;}
.le-card .le-title {text-align:center;color:#f9fafb;font-size:1.45rem;
  letter-spacing:.04em;text-transform:uppercase;margin:0 0 18px;}
.le-card .le-heading {color:#93c5fd;font-size:1.08rem;margin:20px 0 6px;
  border-bottom:1px solid #1f2937;padding-bottom:4px;}
.le-card p {margin:6px 0;text-align:justify;}
.le-card .le-list {margin:4px 0 8px 22px;padding:0;}
.le-card li {margin:3px 0;}
.le-card strong {color:#f9fafb;}
.le-card hr {border:none;border-top:1px solid #374151;margin:14px 0;}
</style>
"""


def today_str() -> str:
    return date.today().strftime("%B %d, %Y")
