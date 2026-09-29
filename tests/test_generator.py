"""Tests for the formatting utilities in ai_core/generator.py."""
import io

import pytest

from ai_core.generator import (format_docx, format_html_preview, format_txt, inline_runs,
                               parse_blocks, safe_filename, sanitize_text, split_terms)

SAMPLE = """```markdown
## Freelance Work Contract

This Agreement is made on “April 15, 2025” — between the parties.

### 1. Services
The **Service Provider** agrees to deliver the work.
- Payment within *7 days*
- Confidentiality at all times
  - Nested point

**2. Payment:**

3. Termination:
Either party may terminate with 15 days notice.

Signature: ______________________
Name: [Name]
```"""


def test_sanitize_removes_smart_quotes_and_fences():
    clean = sanitize_text(SAMPLE)
    assert "“" not in clean and "—" not in clean
    assert '"April 15, 2025"' in clean
    assert not clean.startswith("```") and not clean.endswith("```")
    assert sanitize_text("Fee: ₹5000") == "Fee: Rs. 5000"


def test_parse_blocks_structure():
    kinds = [(b.kind, b.text) for b in parse_blocks(SAMPLE) if b.kind != "blank"]
    assert kinds[0] == ("title", "Freelance Work Contract")
    assert ("heading", "1. Services") in kinds
    assert ("subheading", "2. Payment:") in kinds
    assert ("heading", "3. Termination") in kinds
    bullets = [b for b in parse_blocks(SAMPLE) if b.kind == "bullet"]
    assert len(bullets) == 3 and bullets[2].level == 1


def test_inline_runs_bold_italic_and_signature_lines():
    assert inline_runs("The **Client** pays *now*") == [
        ("The ", False, False), ("Client", True, False), (" pays ", False, False), ("now", False, True)]
    assert inline_runs("Signature: ______________________") == [
        ("Signature: ______________________", False, False)]
    assert inline_runs("file_name_here") == [("file_name_here", False, False)]


def test_split_terms_and_filename():
    assert split_terms("a; b ;; c ") == ["a", "b", "c"]
    assert split_terms("") == []
    assert safe_filename("Non-Disclosure Agreement (NDA)", "pdf") == "non_disclosure_agreement_nda.pdf"


def test_txt_export_has_no_markdown():
    txt = format_txt(SAMPLE)
    assert "**" not in txt and "##" not in txt
    assert txt.startswith("FREELANCE WORK CONTRACT")


def test_html_preview_escapes_html():
    out = format_html_preview("## Title\n<script>alert(1)</script> **bold**")
    assert "<script>" not in out and "&lt;script&gt;" in out
    assert "<strong>bold</strong>" in out and "le-title" in out


def test_docx_is_valid_and_contains_content():
    from docx import Document

    data = format_docx(SAMPLE, "Freelance Work Contract", terms="Pay on time; Keep secrets")
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "FREELANCE WORK CONTRACT" in body
    assert "Service Provider" in body
    assert len(doc.tables) == 1 and doc.tables[0].rows[2].cells[1].text == "Keep secrets"
    assert doc.styles["Normal"].font.name == "Times New Roman"
    assert "All Rights Reserved" in doc.sections[0].footer.paragraphs[0].text


def test_docx_with_custom_logo_bytes():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 80), "red").save(buf, format="JPEG")
    data = format_docx("## T\nHello", "Test", logo=buf.getvalue())
    assert data[:2] == b"PK"  # docx is a zip file


def test_pdf_is_valid():
    pytest.importorskip("fpdf")
    from ai_core.generator import format_pdf

    long_text = SAMPLE + "\n" + "\n".join(f"Clause {i}. " + "Lorem ipsum " * 40 for i in range(30))
    data = format_pdf(long_text, "Freelance Work Contract", terms="Pay on time; Keep secrets")
    assert data.startswith(b"%PDF") and len(data) > 2000
